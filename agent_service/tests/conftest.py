import json
from pathlib import Path

import pytest
from docx import Document
from fastapi.testclient import TestClient

from app.config import get_settings
from app.llm import LLMResult

SECRET = "test-secret"
HEADERS = {"X-Service-Token": SECRET}

ANALYST_JSON = {
    "confidence_score": 4,
    "milestone_status_flagged": "yes",
    "status": "CDR is at risk: radar adapter task is overdue.",
    "explanation": "Reference says On Track but the linked task is overdue.",
    "risks_recommended_for_removal": ["RISK-021"],
}


class FakeLLM:
    """Returns queued replies; records every call."""

    model = "fake-model"

    def __init__(self, replies=None):
        self.replies = list(replies or [])
        self.calls = []

    async def complete(self, *, system, messages, max_tokens):
        self.calls.append(
            dict(system=system, messages=messages, max_tokens=max_tokens)
        )
        reply = self.replies.pop(0)
        if isinstance(reply, Exception):
            raise reply
        # A queued (text, stop_reason) tuple overrides the default "end_turn", e.g. to
        # simulate a reply cut off at max_tokens before any visible text.
        stop_reason = "end_turn"
        if isinstance(reply, tuple):
            reply, stop_reason = reply
        return LLMResult(text=reply, input_tokens=10, output_tokens=5, model="fake-model", stop_reason=stop_reason)


def make_analyst_docx(path: Path) -> None:
    doc = Document()
    doc.add_paragraph("You are an experienced technical program manager.")
    table = doc.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "field"
    table.rows[0].cells[1].text = "meaning"
    doc.add_paragraph("Return JSON only.")
    doc.save(str(path))


@pytest.fixture
def agents_dir(tmp_path):
    d = tmp_path / "agents"
    d.mkdir()
    (d / "milestone_risk_analyst.md").write_text("You are an experienced technical program manager.\nReturn JSON only.")
    (d / "milestone_risk_executor.md").write_text("You are the executor.\nLead with status.")
    (d / "milestone_slide_title.md").write_text("You compress a paragraph into a title and subtitle.")
    return d


@pytest.fixture
def client(monkeypatch, agents_dir):
    monkeypatch.setenv("SERVICE_SHARED_SECRET", SECRET)
    monkeypatch.setenv("ANTHROPIC_API_KEY", "sk-test")
    monkeypatch.setenv("AGENTS_DIR", str(agents_dir))
    monkeypatch.setenv("LOG_LEVEL", "WARNING")
    get_settings.cache_clear()
    from app.main import app

    with TestClient(app) as c:
        c.fake = FakeLLM()
        app.state.llm = c.fake
        yield c
    get_settings.cache_clear()


def j(obj) -> str:
    return json.dumps(obj)
